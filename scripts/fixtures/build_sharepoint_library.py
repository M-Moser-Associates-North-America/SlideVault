"""Generate a synthetic library laid out like M Moser's SharePoint, in tests/fixtures/sharepoint-library/.

    .venv/bin/python scripts/fixtures/build_sharepoint_library.py

Why it exists: the three decks in tests/fixtures/ prove export fidelity, but search, tagging and
folder-path metadata need a bigger library with known answers. Alina described the real archive on
Oct 8 as year folders, then one folder per client or project, plus a "Best of the Best" set sorted by
category (IDD, Business Proposal, SME). This mirrors both layouts:

    sharepoint-library/Submittals/<year>/<client>/<client> - <category>.pptx      (12 decks)
    sharepoint-library/Best of the Best/<category>/<client> - <category>.pptx     (6 copies, 2 per category)

answer_key.csv lists every slide with the slide type planted in it (Cover, About us, Team, Case study,
Values and so on), so AI tagging and search can be scored against known answers.

Every client, person, place and number is invented. No client material. Slide content is chosen with a
fixed seed, so the slides are the same on every run (file bytes differ because zip timestamps change).
The default ingest reads only the top level of tests/fixtures/, so this folder doesn't change the
existing library, export checks or accessibility tests.
"""

from __future__ import annotations

import csv
import random
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

random.seed(7)  # same slide content on every run

OUT = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "sharepoint-library"
ORANGE = RGBColor(0xFF, 0x3F, 0x0A)  # same accent as build_fixtures.py
DARK = RGBColor(0x22, 0x26, 0x2B)
GREY = RGBColor(0x6B, 0x72, 0x7A)
FONT = "Arial"  # installed everywhere, so thumbnails match across machines (as in build_fixtures.py)

# Made-up clients: name, sector, city, what they're building
CLIENTS = [
    ("Northgate Bank", "Financial services", "New York", "a new regional headquarters"),
    ("Lumen Biotech", "Life sciences", "Boston", "a lab and office campus"),
    ("Cobalt Games", "Technology", "Seattle", "a studio for 400 developers"),
    ("Halcyon Health", "Healthcare", "Chicago", "an outpatient clinic network"),
    ("Atlas Logistics", "Logistics", "Dallas", "a hub office and training center"),
    ("Verdant Foods", "Consumer goods", "Toronto", "a test kitchen and workplace"),
    ("Kestrel Law", "Legal", "Washington DC", "a 12-floor law office"),
    ("Orbit Telecom", "Telecommunications", "Denver", "a network operations center"),
    ("Juniper University", "Education", "Austin", "a student innovation hub"),
    ("Sable Hotels", "Hospitality", "Miami", "a flagship hotel lobby and offices"),
    ("Pioneer Energy", "Energy", "Houston", "a low-carbon headquarters"),
    ("Mosaic Media", "Media", "Los Angeles", "a broadcast and content studio"),
]
YEARS = [2024, 2025, 2026]
CATEGORIES = ["IDD", "Business Proposal", "SME"]

TEAM_NAMES = ["Priya Raman", "Marcus Lee", "Elena Ortiz", "Sam Okafor", "Hana Kim",
              "David Brooks", "Aisha Patel", "Tom Fischer", "Lucia Romero", "Ben Carter"]
ROLES = ["Project Director", "Design Lead", "Workplace Strategist", "Engineering Lead",
         "Sustainability Lead", "Project Manager"]


def new_deck(wide=True):
    prs = Presentation()
    if wide:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def add_title_text(prs, slide, title, lines, title_size=32):
    """Title at the top, bullet lines below. Real text boxes, so it stays editable."""
    w = prs.slide_width
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), w - Inches(1.2), Inches(1.0))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size, p.font.bold, p.font.color.rgb = Pt(title_size), True, DARK
    p.alignment = PP_ALIGN.LEFT
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.35), Inches(1.2), Inches(0.08))
    bar.fill.solid(); bar.fill.fore_color.rgb = ORANGE; bar.line.fill.background()
    if lines:
        body = slide.shapes.add_textbox(Inches(0.6), Inches(1.7), w - Inches(1.2), Inches(4.8))
        tf = body.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            para.font.size, para.font.color.rgb = Pt(18), GREY
            para.space_after = Pt(8)
            para.alignment = PP_ALIGN.LEFT


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---- one function per slide type. Each returns the slide type label for the answer key.

def s_cover(prs, c, cat, year):
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.4), prs.slide_width - Inches(1.6), Inches(2.5))
    tf = tb.text_frame
    tf.paragraphs[0].text = f"{c[0]}"
    tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold = Pt(44), True
    tf.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p = tf.add_paragraph(); p.text = f"{cat} for {c[3]} in {c[2]}"
    p.font.size, p.font.color.rgb = Pt(22), ORANGE
    p = tf.add_paragraph(); p.text = f"Submitted {year}"
    p.font.size, p.font.color.rgb = Pt(16), RGBColor(0xCC, 0xCC, 0xCC)
    return "Cover"


def s_about(prs, c, cat, year):
    add_title_text(prs, blank(prs), "About us", [
        "A global design and engineering practice working across workplace, hospitality and education.",
        f"Studios in {random.randint(18, 30)} cities, with teams who design, engineer and deliver under one roof.",
        "We plan, design and build spaces that help organizations work better.",
    ])
    return "About us"


def s_team(prs, c, cat, year):
    s = blank(prs)
    add_title_text(prs, s, f"Your team for {c[0]}", [])
    people = random.sample(TEAM_NAMES, 4)
    roles = random.sample(ROLES, 4)
    for i, (name, role) in enumerate(zip(people, roles)):
        x = Inches(0.6 + i * 3.1)
        circle = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.7), Inches(2.0), Inches(1.4), Inches(1.4))
        circle.fill.solid(); circle.fill.fore_color.rgb = RGBColor(0xE6, 0xE8, 0xEB); circle.line.fill.background()
        tb = s.shapes.add_textbox(x, Inches(3.6), Inches(2.8), Inches(1.2))
        tb.text_frame.paragraphs[0].text = name
        tb.text_frame.paragraphs[0].font.size, tb.text_frame.paragraphs[0].font.bold = Pt(18), True
        p = tb.text_frame.add_paragraph(); p.text = role; p.font.size, p.font.color.rgb = Pt(14), GREY
    return "Team"


def s_client(prs, c, cat, year):
    add_title_text(prs, blank(prs), f"What we heard from {c[0]}", [
        f"{c[0]} is planning {c[3]} in {c[2]}.",
        f"Priorities: attract talent in a competitive {c[1].lower()} market, support hybrid work, and cut energy use.",
        "Success means a space people choose to come to, delivered on time and on budget.",
    ])
    return "About the client"


def s_approach(prs, c, cat, year):
    s = blank(prs)
    add_title_text(prs, s, "Our approach", [])
    steps = ["Discover", "Strategize", "Design", "Deliver"]
    for i, step in enumerate(steps):
        box = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(0.6 + i * 3.1), Inches(2.6), Inches(2.9), Inches(1.2))
        box.fill.solid(); box.fill.fore_color.rgb = ORANGE if i == 0 else RGBColor(0xF3, 0xD2, 0xBC)
        box.line.fill.background()
        box.text_frame.text = step
        box.text_frame.paragraphs[0].font.size, box.text_frame.paragraphs[0].font.bold = Pt(18), True
        box.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else DARK
    return "Approach"


def s_case(prs, c, cat, year):
    other = random.choice([x for x in CLIENTS if x != c])
    add_title_text(prs, blank(prs), f"Case study: {other[0]}", [
        f"Sector: {other[1]}. Location: {other[2]}.",
        f"We delivered {other[3]}, moving {random.randint(300, 1500)} people in {random.randint(9, 20)} months.",
        f"Result: {random.randint(15, 40)}% more space used each day and a LEED Gold rating.",
    ])
    return "Case study"


def s_connect(prs, c, cat, year):
    add_title_text(prs, blank(prs), "Why we're the right partner", [
        f"We've delivered {random.randint(5, 25)} projects for {c[1].lower()} clients.",
        f"Our {c[2]} studio is minutes from your site.",
        "One team from strategy to move-in, so nothing gets lost between phases.",
    ])
    return "How we connect"


def s_values(prs, c, cat, year):
    add_title_text(prs, blank(prs), "Our values", [
        "People first: we design around how people actually work.",
        "Shared ownership: our teams own the outcome with you.",
        "Responsible design: every project has a carbon and wellbeing target.",
    ])
    return "Values"


def s_timeline(prs, c, cat, year):
    s = blank(prs)
    add_title_text(prs, s, "Project timeline", [])
    rows = [("Phase", "Duration", "Key output"),
            ("Strategy", "6 weeks", "Workplace brief"),
            ("Concept design", "8 weeks", "Concept package"),
            ("Detailed design", "12 weeks", "Construction documents"),
            ("Build and move-in", f"{random.randint(20, 40)} weeks", "Handover")]
    table = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(1.8), prs.slide_width - Inches(1.2), Inches(3.5)).table
    for r, row in enumerate(rows):
        for col, val in enumerate(row):
            table.cell(r, col).text = val
    return "Timeline"


def s_fees(prs, c, cat, year):
    s = blank(prs)
    add_title_text(prs, s, "Fee summary by phase", [])
    data = CategoryChartData()
    data.categories = ["Strategy", "Concept", "Detailed design", "Delivery"]
    data.add_series("Fee (USD thousands)", [random.randint(80, 200) for _ in range(4)])
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.8), Inches(8), Inches(4.8), data)
    return "Fees"


def s_sme(prs, c, cat, year):
    topic = random.choice(["Sustainability", "Acoustics", "Lighting design", "Wellbeing", "Change management"])
    add_title_text(prs, blank(prs), f"Expert insight: {topic}", [
        f"Our {topic.lower()} specialists join the project from day one.",
        f"For {c[0]}, that means {topic.lower()} targets are set in the brief, not added at the end.",
    ])
    return "Subject matter expertise"


def s_contact(prs, c, cat, year):
    add_title_text(prs, blank(prs), "Next steps", [
        "We'd welcome the chance to walk you through this proposal in person.",
        f"Contact: {random.choice(TEAM_NAMES)}, {c[2]} studio",
    ])
    return "Next steps"


# Which slides each category of deck gets, in order
PLAN = {
    "Business Proposal": [s_cover, s_about, s_client, s_team, s_approach, s_case, s_connect, s_values, s_timeline, s_fees, s_contact],
    "IDD": [s_cover, s_about, s_client, s_approach, s_case, s_case, s_team, s_timeline, s_contact],
    "SME": [s_cover, s_about, s_sme, s_sme, s_team, s_case, s_connect, s_contact],
}


def use_font(prs):
    """Set every run, including table cells, to FONT."""
    for slide in prs.slides:
        for shape in slide.shapes:
            frames = [shape.text_frame] if shape.has_text_frame else []
            if getattr(shape, "has_table", False) and shape.has_table:
                frames += [c.text_frame for r in shape.table.rows for c in r.cells]
            for tf in frames:
                for para in tf.paragraphs:
                    for run in para.runs:
                        run.font.name = FONT


def build(client, cat, year, wide):
    prs = new_deck(wide)
    types = [fn(prs, client, cat, year) for fn in PLAN[cat]]
    use_font(prs)
    return prs, types


def main():
    rows = []
    # 12 submittals: one per client, spread across years and categories.
    # Two are old 4:3 decks, to test different slide sizes.
    best_count = {cat: 0 for cat in CATEGORIES}
    for i, client in enumerate(CLIENTS):
        year = YEARS[i % 3]
        cat = CATEGORIES[(i // 3) % 3]  # mixes categories across years
        wide = i not in (4, 9)
        prs, types = build(client, cat, year, wide)
        path = OUT / "Submittals" / str(year) / client[0] / f"{client[0]} - {cat}.pptx"
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(path)
        for n, t in enumerate(types, 1):
            rows.append([path.relative_to(OUT), year, client[0], client[1], cat, n, t, "no"])
        # The first two decks of each category also go into Best of the Best
        if best_count[cat] < 2:
            best_count[cat] += 1
            bob = OUT / "Best of the Best" / cat / path.name
            bob.parent.mkdir(parents=True, exist_ok=True)
            prs.save(bob)
            for n, t in enumerate(types, 1):
                rows.append([bob.relative_to(OUT), year, client[0], client[1], cat, n, t, "yes"])

    with open(OUT / "answer_key.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["file", "year", "client", "sector", "category", "slide_number", "slide_type", "best_of_the_best"])
        w.writerows(rows)

    decks = len({r[0] for r in rows})
    print(f"Made {decks} decks and {len(rows)} slide rows in {OUT}")


if __name__ == "__main__":
    main()
