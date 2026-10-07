"""Generate the synthetic fixture decks in tests/fixtures/.

    python3 -m venv .venv && .venv/bin/pip install -r scripts/fixtures/requirements.txt
    .venv/bin/python scripts/fixtures/build_fixtures.py

Every name, number and image here is invented. The decks exist so development, tests,
screenshots and demos never use client material (docs/PRODUCT_BRIEF.md, "Source material").
The output is deterministic, so re-running produces the same slides.
"""

from __future__ import annotations

import struct
import zlib
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parents[2] / "tests" / "fixtures"
INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0xFF, 0x3F, 0x0A)
FONT = "Arial"  # installed everywhere, so renders match across machines


# ---------------------------------------------------------------- helpers


def png(width: int, height: int, seed: int) -> bytes:
    """A small synthetic 'photo': soft gradient with blocks. Pure stdlib, deterministic."""
    rows = []
    for y in range(height):
        row = bytearray([0])  # filter byte
        for x in range(width):
            band = ((x // 40) + (y // 30) + seed) % 4
            base = (70 + band * 30, 90 + band * 25, 110 + band * 20)
            shade = (x + y) // 12
            row += bytes(min(255, c + shade) for c in base)
        rows.append(bytes(row))
    raw = zlib.compress(b"".join(rows), 9)

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", raw) + chunk(b"IEND", b"")


def text_box(slide, left, top, width, height, lines, size=18, bold_first=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = para.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = INK
        run.font.bold = bold_first and i == 0
    return box


def title(slide, text, prs, size=36):
    return text_box(slide, Inches(0.6), Inches(0.4), prs.slide_width - Inches(1.2), Inches(1.2), [text], size=size, bold_first=True)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # "Blank" layout in the default template


def image(slide, prs, path: Path, left, top, width):
    return slide.shapes.add_picture(str(path), left, top, width=width)


# ---------------------------------------------------------------- decks


def northwind_capabilities(img: Path) -> Presentation:
    """Canonical 'about us' content for a fictional workplace design studio (16:9)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = blank(prs)
    title(s, "Northwind Studio", prs, size=54)
    text_box(s, Inches(0.6), Inches(1.8), Inches(10), Inches(1), ["Workplace strategy and interior design since 1998"], size=24)

    s = blank(prs)
    title(s, "Who we are", prs)
    text_box(s, Inches(0.6), Inches(1.6), Inches(7), Inches(4.5), [
        "An independent studio of 140 designers, strategists and project managers.",
        "Offices in Lisbon, Toronto and Singapore.",
        "We design places where teams choose to work together.",
    ])
    image(s, prs, img, Inches(8.2), Inches(1.6), Inches(4.4))
    notes(s, "Speaker note: mention the 25-year anniversary only for long-standing clients.")

    s = blank(prs)
    title(s, "Services", prs)
    text_box(s, Inches(0.6), Inches(1.6), Inches(12), Inches(4.5), [
        "Workplace strategy and occupancy studies",
        "Concept and detailed interior design",
        "Change management for hybrid teams",
        "Project and cost management through handover",
    ], size=22)

    s = blank(prs)
    title(s, "Leadership team", prs)
    rows = [("Name", "Role", "Based in"),
            ("Ines Calder", "Managing principal", "Lisbon"),
            ("Tomas Reyna", "Design director", "Toronto"),
            ("Priya Anand", "Strategy lead", "Singapore")]
    table = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(1.6), Inches(12), Inches(2.4)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.text_frame.paragraphs[0].runs[0].font.name = FONT
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(16)

    s = blank(prs)
    title(s, "Internal: pricing guidance", prs)
    text_box(s, Inches(0.6), Inches(1.6), Inches(12), Inches(3), [
        "Hidden slide. Not for client decks.",
        "Day rates and discount bands are agreed per pursuit.",
    ])
    s._element.set("show", "0")  # hidden in the source deck
    notes(s, "Speaker note: never export this slide without approval.")

    s = blank(prs)
    title(s, "Our approach in four phases", prs)
    for i, label in enumerate(["Listen", "Frame", "Design", "Deliver"]):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6 + i * 3.1), Inches(2.4), Inches(2.8), Inches(1.6))
        shape.fill.solid()
        shape.fill.fore_color.rgb = ACCENT if i == 0 else RGBColor(0xEC, 0xEE, 0xF1)
        shape.line.color.rgb = INK
        shape.text_frame.text = f"{i + 1} · {label}"
        run = shape.text_frame.paragraphs[0].runs[0]
        run.font.name, run.font.size, run.font.color.rgb = FONT, Pt(20), INK
    return prs


def harbor_bank_proposal(img: Path) -> Presentation:
    """A reusable pursuit deck for a fictional bank client (16:9), with a native chart."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = blank(prs)
    title(s, "Harbor Bank: headquarters consolidation", prs, size=44)
    text_box(s, Inches(0.6), Inches(1.9), Inches(10), Inches(1), ["Proposal prepared by Northwind Studio"], size=22)

    s = blank(prs)
    title(s, "The challenge", prs)
    text_box(s, Inches(0.6), Inches(1.6), Inches(12), Inches(4), [
        "Three leased buildings with peak desk use of 41%.",
        "Teams split across sites lose time to travel and booking.",
        "Leases on two buildings end within 30 months.",
    ])

    s = blank(prs)
    title(s, "Utilisation evidence", prs)
    data = CategoryChartData()
    data.categories = ["Mon", "Tue", "Wed", "Thu", "Fri"]
    data.add_series("Peak desk use (%)", (33, 41, 39, 36, 18))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.6), Inches(8), Inches(5), data)
    text_box(s, Inches(9), Inches(2), Inches(3.8), Inches(3), ["Sensor study, six weeks, 1,200 desks."], size=16)
    notes(s, "Speaker note: source is the synthetic occupancy study in the appendix.")

    s = blank(prs)
    title(s, "Recommended concept", prs)
    image(s, prs, img, Inches(0.6), Inches(1.6), Inches(6))
    text_box(s, Inches(7), Inches(1.6), Inches(5.8), Inches(4), [
        "One address, organised around team neighbourhoods.",
        "Shared amenity floor at level 3.",
        "A 1:1.3 sharing ratio validated by the study.",
    ])
    return prs


def legacy_case_study(img: Path) -> Presentation:
    """An older 4:3 case-study deck, to cover mixed slide sizes."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9144000), Emu(6858000)  # 10 x 7.5 in (4:3)

    s = blank(prs)
    title(s, "Case study: Juniper Labs research campus", prs, size=32)
    image(s, prs, img, Inches(0.6), Inches(1.8), Inches(5))
    text_box(s, Inches(6), Inches(1.8), Inches(3.6), Inches(4), ["Completed 2019.", "4 buildings, 22,000 m².", "Lab-to-office ratio 60:40."], size=16)

    s = blank(prs)
    title(s, "Outcome", prs, size=32)
    text_box(s, Inches(0.6), Inches(1.8), Inches(8.8), Inches(4), [
        "Collaboration space use up 2.4x in the first year.",
        "Delivered two weeks ahead of programme.",
    ], size=20)
    return prs


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    photo = OUT / "_synthetic-photo.png"
    photo.write_bytes(png(480, 300, seed=3))
    decks = {
        "northwind-capabilities.pptx": ("Northwind Studio capabilities", northwind_capabilities(photo)),
        "harbor-bank-proposal.pptx": ("Harbor Bank HQ proposal", harbor_bank_proposal(photo)),
        "legacy-case-study-4x3.pptx": ("Juniper Labs case study (4:3)", legacy_case_study(photo)),
    }
    for name, (deck_title, prs) in decks.items():
        prs.core_properties.author = "SlideVault fixtures"
        prs.core_properties.title = deck_title
        prs.save(OUT / name)
        print(f"wrote tests/fixtures/{name}: {len(prs.slides)} slides")
    photo.unlink()


if __name__ == "__main__":
    main()
